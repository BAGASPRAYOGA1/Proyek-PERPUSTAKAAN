from tkinter import ttk
import tkinter as tk
from tkinter import messagebox
from datetime import datetime
import os
import sys

import database

# =========================
# ReportLab
# =========================
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

# =========================
# Openpyxl
# =========================
from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)

FONT_FAMILY = "Segoe UI"


class RekapBukuFrame(ttk.Frame):

    def __init__(self, master):
        super().__init__(master)

        self.pack(fill="both", expand=True)

        self.configure_style()

        self.create_widgets()

        self.load_data()

    # ==================================================

    def configure_style(self):

        style = ttk.Style(self)

        style.configure(
            "Rekap.TFrame",
            background="#f0f4f8",
        )

        style.configure(
            "RekapHeader.TLabel",
            background="#f0f4f8",
            foreground="#2c3e50",
            font=(FONT_FAMILY, 18, "bold"),
        )

        style.configure(
            "RekapLabel.TLabel",
            background="#f0f4f8",
            foreground="#34495e",
            font=(FONT_FAMILY, 10),
        )

        style.configure(
            "RekapButton.TButton",
            font=(FONT_FAMILY, 10, "bold"),
            padding=(12, 6),
        )

    # ==================================================

    def create_widgets(self):

        header = ttk.Label(
            self,
            text="Rekap Buku Perpustakaan",
            style="RekapHeader.TLabel",
        )

        header.pack(
            pady=(20, 15)
        )

        # ==========================================
        # PANEL FILTER
        # ==========================================

        filter_panel = ttk.LabelFrame(
            self,
            text=" Filter Rekap Buku ",
            style="Rekap.TFrame",
        )

        filter_panel.pack(
            fill="x",
            padx=30,
            pady=(0, 10),
        )

        lbl_judul = ttk.Label(
            filter_panel,
            text="Judul Buku :",
            style="RekapLabel.TLabel",
        )

        lbl_judul.grid(
            row=0,
            column=0,
            padx=(15, 10),
            pady=8,
            sticky="w",
        )

        self.entry_cari = ttk.Entry(
            filter_panel,
            width=35,
            font=(FONT_FAMILY, 10),
        )

        self.entry_cari.grid(
            row=0,
            column=1,
            padx=(0, 20),
            pady=8,
            sticky="w",
        )

        lbl_kelas = ttk.Label(
            filter_panel,
            text="Kelas :",
            style="RekapLabel.TLabel",
        )

        lbl_kelas.grid(
            row=0,
            column=2,
            padx=(10, 10),
            pady=8,
            sticky="w",
        )

        self.combo_kelas = ttk.Combobox(
            filter_panel,
            values=[
                "Semua",
                "1",
                "2",
                "3",
                "4",
                "5",
                "6",
            ],
            state="readonly",
            width=12,
            font=(FONT_FAMILY, 10),
        )

        self.combo_kelas.set("Semua")

        self.combo_kelas.grid(
            row=0,
            column=3,
            padx=(0, 20),
            pady=8,
            sticky="w",
        )

        filter_panel.columnconfigure(5, weight=1)

        # ==========================================
        # PANEL TOMBOL
        # ==========================================

        btn_frame = ttk.Frame(
            self,
            style="Rekap.TFrame",
        )

        btn_frame.pack(
            fill="x",
            padx=30,
            pady=(0, 10),
        )

        tombol = [

            ("Tampilkan", "tampilkan"),

            ("Export Excel", "excel"),

            ("Export PDF", "pdf"),

            ("Refresh", "refresh"),

        ]

        self.buttons = {}

        for i, (text, key) in enumerate(tombol):

            btn = ttk.Button(
                btn_frame,
                text=text,
                style="RekapButton.TButton",
                command=lambda k=key: self.on_button_click(k),
            )

            btn.grid(
                row=0,
                column=i,
                padx=3,
                pady=2,
                sticky="ew",
            )

            self.buttons[key] = btn

        btn_frame.columnconfigure((0, 1, 2, 3), weight=1)

                # ==================================================
        # TREEVIEW
        # ==================================================

        tree_frame = ttk.Frame(
            self,
            style="Rekap.TFrame",
        )

        tree_frame.pack(
            fill="both",
            expand=True,
            padx=30,
            pady=(0, 15),
        )

        tree_frame.columnconfigure(0, weight=1)
        tree_frame.rowconfigure(0, weight=1)

        columns = (
            "kode_buku",
            "judul",
            "kelas",
            "stok_total",
            "dipinjam",
            "stok_tersedia",
        )

        self.tree = ttk.Treeview(
            tree_frame,
            columns=columns,
            show="headings",
            selectmode="browse",
        )

        self.tree.heading("kode_buku", text="Kode Buku")
        self.tree.heading("judul", text="Judul Buku")
        self.tree.heading("kelas", text="Kelas")
        self.tree.heading("stok_total", text="Total")
        self.tree.heading("dipinjam", text="Dipinjam")
        self.tree.heading("stok_tersedia", text="Tersedia")

        self.tree.column("kode_buku", width=120, anchor="center")
        self.tree.column("judul", width=330, anchor="w")
        self.tree.column("kelas", width=70, anchor="center")
        self.tree.column("stok_total", width=90, anchor="center")
        self.tree.column("dipinjam", width=90, anchor="center")
        self.tree.column("stok_tersedia", width=90, anchor="center")

        self.tree.grid(
            row=0,
            column=0,
            sticky="nsew",
        )

        scroll_y = ttk.Scrollbar(
            tree_frame,
            orient="vertical",
            command=self.tree.yview,
        )

        scroll_y.grid(
            row=0,
            column=1,
            sticky="ns",
        )

        self.tree.configure(
            yscrollcommand=scroll_y.set
        )

        scroll_x = ttk.Scrollbar(
            tree_frame,
            orient="horizontal",
            command=self.tree.xview,
        )

        scroll_x.grid(
            row=1,
            column=0,
            sticky="ew",
        )

        self.tree.configure(
            xscrollcommand=scroll_x.set)

    # ==================================================
    # LOAD DATA
    # ==================================================

    def load_data(self):

        try:

            database.create_tables()

            keyword = self.entry_cari.get().strip()

            kelas = self.combo_kelas.get().strip()

            query = """
                SELECT
                    id_buku,
                    kode_buku,
                    judul,
                    kelas,
                    stok_total,
                    stok_tersedia
                FROM buku
            """

            kondisi = []

            params = []

            if keyword:

                kondisi.append("judul LIKE ?")

                params.append(f"%{keyword}%")

            if kelas != "Semua":

                kondisi.append("kelas=?")

                params.append(kelas)

            if kondisi:

                query += " WHERE "

                query += " AND ".join(kondisi)

            query += """

                ORDER BY

                kelas ASC,

                judul ASC

            """

            rows = database.fetch_all(

                query,

                tuple(params)

            )

            for item in self.tree.get_children():

                self.tree.delete(item)

            if rows:

                for row in rows:

                    stok_total = int(row["stok_total"] or 0)

                    stok_tersedia = int(row["stok_tersedia"] or 0)

                    dipinjam = stok_total - stok_tersedia

                    self.tree.insert(

                        "",

                        "end",

                        iid=row["id_buku"],

                        values=(

                            row["kode_buku"],

                            row["judul"],

                            row["kelas"],

                            stok_total,

                            dipinjam,

                            stok_tersedia,

                        ),

                    )

        except Exception as e:

            print("[ERROR]", e)

            messagebox.showerror(

                "Error",

                str(e)

            )

                # ==================================================
    # BASE DIRECTORY
    # ==================================================

    def _get_base_dir(self):

        if getattr(sys, "frozen", False):

            return os.path.dirname(
                os.path.abspath(sys.executable)
            )

        return os.path.dirname(
            os.path.abspath(sys.argv[0])
        )

    # ==================================================
    # EXPORT FOLDER
    # ==================================================

    def _get_export_folder(self, subfolder):

        base_dir = self._get_base_dir()

        folder = os.path.join(
            base_dir,
            "exports",
            subfolder,
        )

        os.makedirs(
            folder,
            exist_ok=True,
        )

        return folder

    # ==================================================
    # OPEN EXPORT FOLDER
    # ==================================================

    def _open_export_folder(self, folder):

        try:

            os.startfile(folder)

        except Exception as e:

            print(e)

            messagebox.showerror(
                "Error",
                f"Gagal membuka folder.\n\n{e}"
            )

    # ==================================================
    # EXPORT SUCCESS
    # ==================================================

    def _show_export_success(self, file_path):

        try:

            folder = os.path.dirname(file_path)

            jawab = messagebox.askyesno(

                "Export Berhasil",

                f"File berhasil disimpan.\n\n"
                f"{file_path}\n\n"
                "Buka folder hasil export?"

            )

            if jawab:

                self._open_export_folder(folder)

        except Exception as e:

            print(e)

    # ==================================================
    # TAMPILKAN
    # ==================================================

    def tampilkan(self):

        self.load_data()

    # ==================================================
    # REFRESH
    # ==================================================

    def refresh(self):

        self.entry_cari.delete(
            0,
            tk.END,
        )

        self.combo_kelas.set(
            "Semua"
        )

        self.load_data()

    # ==================================================
    # TOTAL BUKU
    # ==================================================

    def get_total_buku(self):

        try:

            query = """

                SELECT
                    COUNT(*) AS total
                FROM buku

            """

            hasil = database.fetch_one(query)

            if hasil:

                return hasil["total"]

            return 0

        except:

            return 0

    # ==================================================
    # TOTAL STOK
    # ==================================================

    def get_total_stok(self):

        try:

            query = """

                SELECT
                    SUM(stok_total) AS total
                FROM buku

            """

            hasil = database.fetch_one(query)

            if hasil["total"] is None:

                return 0

            return hasil["total"]

        except:

            return 0

            # ==================================================
    # EXPORT EXCEL
    # ==================================================

    def export_excel(self):

        try:

            self.load_data()

            items = self.tree.get_children()

            if not items:

                messagebox.showwarning(
                    "Peringatan",
                    "Tidak ada data buku."
                )

                return

            folder = self._get_export_folder("excel")

            nama_file = os.path.join(

                folder,

                f"Rekap_Buku_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.xlsx"

            )

            wb = Workbook()

            ws = wb.active

            ws.title = "Rekap Buku"

            header = [

                "Kode Buku",

                "Judul Buku",

                "Kelas",

                "Stok Total",

                "Sedang Dipinjam",

                "Stok Tersedia",

            ]

            header_font = Font(

                bold=True,

                color="FFFFFF"

            )

            header_fill = PatternFill(

                start_color="2C3E50",

                end_color="2C3E50",

                fill_type="solid",

            )

            center = Alignment(

                horizontal="center",

                vertical="center",

            )

            border = Border(

                left=Side(style="thin"),

                right=Side(style="thin"),

                top=Side(style="thin"),

                bottom=Side(style="thin"),

            )

            for col, text in enumerate(

                header,

                start=1,

            ):

                cell = ws.cell(

                    row=1,

                    column=col,

                    value=text,

                )

                cell.font = header_font

                cell.fill = header_fill

                cell.alignment = center

                cell.border = border

            row_excel = 2

            for item in items:

                values = self.tree.item(

                    item,

                    "values",

                )

                for col, value in enumerate(

                    values,

                    start=1,

                ):

                    cell = ws.cell(

                        row=row_excel,

                        column=col,

                        value=value,

                    )

                    cell.border = border

                    if col == 2:

                        cell.alignment = Alignment(

                            horizontal="left",

                            vertical="center",

                        )

                    else:

                        cell.alignment = center

                row_excel += 1

            ukuran = {

                "A":18,

                "B":45,

                "C":10,

                "D":15,

                "E":18,

                "F":18,

            }

            for kolom, lebar in ukuran.items():

                ws.column_dimensions[kolom].width = lebar

            wb.save(nama_file)

            self._show_export_success(

                nama_file

            )

        except Exception as e:

            print(

                "[ERROR]",

                e,

            )

            messagebox.showerror(

                "Error",

                str(e),

            )

                # ==================================================
    # EXPORT PDF
    # ==================================================

    def export_pdf(self):

        try:

            self.load_data()

            items = self.tree.get_children()

            if not items:

                messagebox.showwarning(

                    "Peringatan",

                    "Tidak ada data buku."

                )

                return

            folder = self._get_export_folder("pdf")

            file_path = os.path.join(

                folder,

                f"Rekap_Buku_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.pdf"

            )

            doc = SimpleDocTemplate(

                file_path,

                pagesize=landscape(A4),

                leftMargin=15 * mm,

                rightMargin=15 * mm,

                topMargin=15 * mm,

                bottomMargin=15 * mm,

            )

            styles = getSampleStyleSheet()

            elements = []

            title_style = ParagraphStyle(

                "Title",

                parent=styles["Title"],

                alignment=1,

                fontName="Helvetica-Bold",

                fontSize=18,

                textColor=colors.HexColor("#2c3e50"),

                spaceAfter=10,

            )

            date_style = ParagraphStyle(

                "Date",

                parent=styles["Normal"],

                alignment=1,

                fontSize=10,

                textColor=colors.HexColor("#7f8c8d"),

                spaceAfter=12,

            )

            elements.append(

                Paragraph(

                    "Rekap Koleksi Buku Perpustakaan",

                    title_style,

                )

            )

            elements.append(

                Paragraph(

                    f"Tanggal Cetak : {datetime.now().strftime('%d-%m-%Y %H:%M:%S')}",

                    date_style,

                )

            )

            elements.append(

                Spacer(

                    1,

                    5 * mm,

                )

            )

            data = [[

                "Kode Buku",

                "Judul Buku",

                "Kelas",

                "Stok",

                "Dipinjam",

                "Tersedia",

            ]]

            for item in items:

                values = self.tree.item(

                    item,

                    "values",

                )

                data.append(

                    list(values)

                )

            table = Table(

                data,

                colWidths=[

                    35 * mm,

                    90 * mm,

                    20 * mm,

                    25 * mm,

                    30 * mm,

                    30 * mm,

                ],

                repeatRows=1,

            )

            table.setStyle(

                TableStyle([

                    ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#2c3e50")),

                    ("TEXTCOLOR",(0,0),(-1,0),colors.white),

                    ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),

                    ("FONTSIZE",(0,0),(-1,0),10),

                    ("GRID",(0,0),(-1,-1),0.5,colors.grey),

                    ("ALIGN",(0,0),(-1,-1),"CENTER"),

                    ("VALIGN",(0,0),(-1,-1),"MIDDLE"),

                    ("ROWBACKGROUNDS",(0,1),(-1,-1),

                        [colors.white, colors.HexColor("#ecf0f1")]),

                    ("TOPPADDING",(0,0),(-1,-1),5),

                    ("BOTTOMPADDING",(0,0),(-1,-1),5),

                ])

            )

            elements.append(

                table

            )

            doc.build(

                elements

            )

            self._show_export_success(

                file_path

            )

        except Exception as e:

            print(e)

            messagebox.showerror(

                "Error",

                str(e)

            )

    # ==================================================
    # BUTTON ACTION
    # ==================================================

    def on_button_click(self, action):

        if action == "tampilkan":

            self.tampilkan()

        elif action == "excel":

            self.export_excel()

        elif action == "pdf":

            self.export_pdf()

        elif action == "refresh":

            self.refresh()

        else:

            messagebox.showinfo(

                "Informasi",

                "Menu belum tersedia."

            )