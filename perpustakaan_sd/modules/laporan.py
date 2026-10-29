from tkinter import ttk
import tkinter as tk
from tkinter import messagebox
from datetime import date, datetime
import os
import sys

import database

# ReportLab imports (export PDF)
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

# openpyxl imports (export Excel)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Font Segoe UI
FONT_FAMILY = "Segoe UI"


class LaporanFrame(ttk.Frame):
    def __init__(self, master):
        super().__init__(master)
        self.pack(fill="both", expand=True)
        self.configure_style()
        self.create_widgets()
        # Tampilkan data laporan langsung saat halaman dibuka
        self.load_data()

    def configure_style(self):
        """Konfigurasi style untuk halaman Laporan."""
        style = ttk.Style(self)
        style.configure("Laporan.TFrame", background="#f0f4f8")
        style.configure(
            "LaporanHeader.TLabel",
            background="#f0f4f8",
            font=(FONT_FAMILY, 18, "bold"),
            foreground="#2c3e50",
        )
        style.configure(
            "LaporanFormLabel.TLabel",
            background="#f0f4f8",
            font=(FONT_FAMILY, 10),
            foreground="#34495e",
        )
        style.configure(
            "LaporanAction.TButton",
            font=(FONT_FAMILY, 10, "bold"),
            padding=(12, 6),
        )

    def create_widgets(self):
        """Membuat layout halaman Laporan."""
        # ===== Header =====
        header = ttk.Label(self, text="Laporan Peminjaman", style="LaporanHeader.TLabel")
        header.pack(pady=(20, 15))

        # ===== Panel Filter =====
        filter_panel = ttk.LabelFrame(
            self, text=" Filter Laporan ", style="Laporan.TFrame"
        )
        filter_panel.pack(fill="x", padx=30, pady=(0, 10))

        # Tahun (Combobox)
        lbl_tahun = ttk.Label(
            filter_panel, text="Tahun :", style="LaporanFormLabel.TLabel"
        )
        lbl_tahun.grid(row=0, column=0, sticky="w", padx=(15, 10), pady=8)

        current_year = date.today().year
        tahun_options = ["Semua"] + [
            str(y) for y in range(current_year, current_year - 10, -1)
        ]
        self.combo_tahun = ttk.Combobox(
            filter_panel,
            values=tahun_options,
            font=(FONT_FAMILY, 10),
            state="readonly",
            width=12,
        )
        self.combo_tahun.set("Semua")
        self.combo_tahun.grid(row=0, column=1, sticky="w", padx=(0, 20), pady=8)

        # Bulan (Combobox)
        lbl_bulan = ttk.Label(
            filter_panel, text="Bulan :", style="LaporanFormLabel.TLabel"
        )
        lbl_bulan.grid(row=0, column=2, sticky="w", padx=(15, 10), pady=8)

        bulan_list = [
            ("Semua", ""),
            ("Januari", "01"),
            ("Februari", "02"),
            ("Maret", "03"),
            ("April", "04"),
            ("Mei", "05"),
            ("Juni", "06"),
            ("Juli", "07"),
            ("Agustus", "08"),
            ("September", "09"),
            ("Oktober", "10"),
            ("November", "11"),
            ("Desember", "12"),
        ]
        self.bulan_labels = [b[0] for b in bulan_list]
        self.bulan_map = {b[0]: b[1] for b in bulan_list}
        self.combo_bulan = ttk.Combobox(
            filter_panel,
            values=self.bulan_labels,
            font=(FONT_FAMILY, 10),
            state="readonly",
            width=14,
        )
        self.combo_bulan.set("Semua")
        self.combo_bulan.grid(row=0, column=3, sticky="w", padx=(0, 20), pady=8)

        # Kelas (Combobox)
        lbl_kelas = ttk.Label(
            filter_panel, text="Kelas :", style="LaporanFormLabel.TLabel"
        )
        lbl_kelas.grid(row=0, column=4, sticky="w", padx=(15, 10), pady=8)

        self.combo_kelas = ttk.Combobox(
            filter_panel,
            values=["Semua", "1", "2", "3", "4", "5", "6"],
            font=(FONT_FAMILY, 10),
            state="readonly",
            width=12,
        )
        self.combo_kelas.set("Semua")
        self.combo_kelas.grid(row=0, column=5, sticky="w", padx=(0, 15), pady=8)

        # Checkbutton: hanya buku yang belum dikembalikan
        self.belum_kembali_var = tk.BooleanVar(value=False)
        self.chk_belum_kembali = ttk.Checkbutton(
            filter_panel,
            text="Tampilkan hanya buku yang belum dikembalikan",
            variable=self.belum_kembali_var,
            style="LaporanFormLabel.TLabel",
        )
        self.chk_belum_kembali.grid(row=0, column=6, sticky="w", padx=(10, 15), pady=8)

        filter_panel.columnconfigure(7, weight=1)

        # ===== Panel Tombol =====
        btn_frame = ttk.Frame(self, style="Laporan.TFrame")
        btn_frame.pack(fill="x", padx=30, pady=(0, 10))

        btn_frame.columnconfigure((0, 1, 2, 3), weight=1)

        buttons = [
            ("Tampilkan", "tampilkan"),
            ("Export Excel", "export_excel"),
            ("Export PDF", "export_pdf"),
            ("Refresh", "refresh"),
        ]

        self.buttons = {}

        for i, (text, key) in enumerate(buttons):
            btn = ttk.Button(
                btn_frame,
                text=text,
                style="LaporanAction.TButton",
                command=lambda k=key: self.on_button_click(k),
            )
            btn.grid(row=0, column=i, padx=3, pady=2, sticky="ew")
            self.buttons[key] = btn

        # ===== Treeview =====
        tree_frame = ttk.Frame(self, style="Laporan.TFrame")
        tree_frame.pack(fill="both", expand=True, padx=30, pady=(0, 15))
        tree_frame.columnconfigure(0, weight=1)
        tree_frame.rowconfigure(0, weight=1)

        columns = (
            "id_pinjam",
            "nama_siswa",
            "kelas",
            "judul_buku",
            "tanggal_pinjam",
            "tanggal_kembali",
            "status",
        )

        self.tree = ttk.Treeview(
            tree_frame,
            columns=columns,
            show="headings",
            selectmode="browse",
        )

        headings = {
            "id_pinjam": "ID",
            "nama_siswa": "Nama Siswa",
            "kelas": "Kelas",
            "judul_buku": "Buku",
            "tanggal_pinjam": "Tanggal Pinjam",
            "tanggal_kembali": "Tanggal Kembali",
            "status": "Status",
        }

        widths = {
            "id_pinjam": 40,
            "nama_siswa": 180,
            "kelas": 60,
            "judul_buku": 220,
            "tanggal_pinjam": 120,
            "tanggal_kembali": 120,
            "status": 110,
        }

        for col in columns:
            self.tree.heading(col, text=headings[col])
            self.tree.column(
                col,
                width=widths[col],
                anchor="center"
                if col in (
                    "id_pinjam",
                    "kelas",
                    "tanggal_pinjam",
                    "tanggal_kembali",
                    "status",
                )
                else "w",
            )

        self.tree.grid(row=0, column=0, sticky="nsew")

        # Scrollbar
        scroll_y = ttk.Scrollbar(
            tree_frame, orient="vertical", command=self.tree.yview
        )
        scroll_y.grid(row=0, column=1, sticky="ns")
        self.tree.configure(yscrollcommand=scroll_y.set)

        scroll_x = ttk.Scrollbar(
            tree_frame, orient="horizontal", command=self.tree.xview
        )
        scroll_x.grid(row=1, column=0, sticky="ew")
        self.tree.configure(xscrollcommand=scroll_x.set)

    def load_data(self):
        """Muat data laporan ke dalam Treeview.

        Filter yang didukung: Tahun, Bulan, Kelas, dan status
        "belum dikembalikan" (status = 'Dipinjam'). Semua nilai
        menggunakan parameter query (?). Jika tidak ada filter yang
        aktif, seluruh data ditampilkan.
        """
        try:
            # Pastikan tabel tersedia
            database.create_tables()

            # Ambil nilai filter
            tahun = self.combo_tahun.get().strip()
            bulan_label = self.combo_bulan.get().strip()
            kelas = self.combo_kelas.get().strip()
            belum_kembali = self.belum_kembali_var.get()

            # Bangun query dasar dengan JOIN
            query = (
                "SELECT p.id_pinjam, p.id_buku, p.id_siswa, "
                "s.nama AS nama_siswa, s.kelas, "
                "b.judul AS judul_buku, "
                "p.tanggal_pinjam, p.tanggal_kembali, p.status "
                "FROM peminjaman p "
                "LEFT JOIN siswa s ON s.id_siswa = p.id_siswa "
                "LEFT JOIN buku b ON b.id_buku = p.id_buku "
            )

            conditions = []
            params = []

            # Filter Tahun
            if tahun and tahun != "Semua":
                conditions.append("substr(p.tanggal_pinjam, 1, 4) = ?")
                params.append(tahun)

            # Filter Bulan
            bulan_code = self.bulan_map.get(bulan_label, "")
            if bulan_label and bulan_label != "Semua" and bulan_code:
                conditions.append("substr(p.tanggal_pinjam, 6, 2) = ?")
                params.append(bulan_code)

            # Filter Kelas
            if kelas and kelas != "Semua":
                conditions.append("s.kelas = ?")
                params.append(kelas)

            # Filter status: hanya buku yang belum dikembalikan
            if belum_kembali:
                conditions.append("p.status = ?")
                params.append("Dipinjam")

            if conditions:
                query += "WHERE " + " AND ".join(conditions) + " "

            query += "ORDER BY p.id_pinjam DESC"

            # Ambil data
            rows = database.fetch_all(query, tuple(params))

            # Hapus isi Treeview lama
            for item in self.tree.get_children():
                self.tree.delete(item)

            # Tampilkan data
            if rows:
                for row in rows:
                    self.tree.insert(
                        "",
                        "end",
                        iid=str(row["id_pinjam"]),
                        values=(
                            row["id_pinjam"],
                            row["nama_siswa"],
                            row["kelas"],
                            row["judul_buku"],
                            row["tanggal_pinjam"],
                            row["tanggal_kembali"],
                            row["status"],
                        ),
                    )
        except Exception as e:
            print(f"[ERROR] Gagal memuat data laporan: {e}")
            messagebox.showerror("Error", f"Terjadi kesalahan: {e}")

    def tampilkan_laporan(self):
        """Tampilkan laporan berdasarkan filter yang dipilih."""
        try:
            self.load_data()
            messagebox.showinfo("Informasi", "Laporan berhasil ditampilkan.")
        except Exception as e:
            print(f"[ERROR] Gagal menampilkan laporan: {e}")
            messagebox.showerror("Error", f"Terjadi kesalahan: {e}")

    def _get_base_dir(self):
        """Kembalikan folder dasar untuk menyimpan file export.

        Berfungsi baik saat dijalankan dari Python maupun sebagai
        file .exe hasil build PyInstaller (sys.frozen).
        """
        if getattr(sys, "frozen", False):
            # Saat di-build menjadi .exe, gunakan folder tempat exe berada
            return os.path.dirname(os.path.abspath(sys.executable))
        # Saat dijalankan dari Python, gunakan folder tempat script berada
        return os.path.dirname(os.path.abspath(sys.argv[0]))

    def _get_export_folder(self, subfolder):
        """Kembalikan path folder export dan buat otomatis jika belum ada.

        Folder dibuat dengan os.makedirs(path, exist_ok=True).
        """
        base_dir = self._get_base_dir()
        folder_path = os.path.join(base_dir, "exports", subfolder)
        os.makedirs(folder_path, exist_ok=True)
        return folder_path

    def _open_export_folder(self, folder_path):
        """Buka folder hasil export menggunakan Windows File Explorer.

        Jika gagal, tampilkan MessageBox Error.
        """
        try:
            os.startfile(folder_path)
        except Exception as e:
            print(f"[ERROR] Gagal membuka folder export: {e}")
            messagebox.showerror(
                "Error",
                f"Gagal membuka folder hasil export:\n{folder_path}\n\n{e}",
            )

    def _show_export_success(self, file_path):
        """Tampilkan MessageBox Yes/No setelah export berhasil.

        Jika YES, buka folder tempat file berada. Jika NO, tutup saja.
        """
        try:
            folder_path = os.path.dirname(file_path)
            jawaban = messagebox.askyesno(
                "Export Berhasil",
                f"File berhasil disimpan.\n\n"
                f"Lokasi:\n{file_path}\n\n"
                f"Apakah ingin membuka folder hasil export?",
            )
            if jawaban:
                self._open_export_folder(folder_path)
        except Exception as e:
            print(f"[ERROR] Gagal menampilkan dialog export: {e}")
            messagebox.showerror("Error", f"Terjadi kesalahan: {e}")

    def export_excel(self):
        """Export laporan ke file Excel menggunakan openpyxl.

        Membuat file Laporan_Perpustakaan_<tanggal-jam>.xlsx berisi
        header dan seluruh data transaksi yang tampil di Treeview
        (mengikuti filter yang aktif). Menyimpan ke folder
        exports/excel/ lalu menampilkan dialog berhasil.
        """
        try:
            # Pastikan data laporan sudah dimuat sesuai filter
            self.load_data()

            # Ambil data dari Treeview (mengikuti filter yang aktif)
            items = self.tree.get_children()
            if not items:
                messagebox.showwarning(
                    "Peringatan",
                    "Tidak ada data transaksi untuk diekspor.",
                )
                return

            # Folder tujuan dan nama file unik (tanggal & jam)
            folder_path = self._get_export_folder("excel")
            timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
            file_name = f"Laporan_Perpustakaan_{timestamp}.xlsx"
            file_path = os.path.join(folder_path, file_name)

            # Buat workbook baru
            wb = Workbook()
            ws = wb.active
            ws.title = "Laporan Peminjaman"

            # Header
            header = [
                "ID",
                "Nama Siswa",
                "Kelas",
                "Buku",
                "Tanggal Pinjam",
                "Tanggal Kembali",
                "Status",
            ]

            # Gaya header
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(
                start_color="2C3E50", end_color="2C3E50", fill_type="solid"
            )
            header_align = Alignment(horizontal="center", vertical="center")
            thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

            # Tulis header
            for col_idx, title in enumerate(header, start=1):
                cell = ws.cell(row=1, column=col_idx, value=title)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
                cell.border = thin_border

            # Tulis data dari Treeview
            for row_idx, item in enumerate(items, start=2):
                values = self.tree.item(item, "values")
                for col_idx, value in enumerate(values, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.border = thin_border
                    # Rata tengah untuk kolom ID, Kelas, Tanggal, Status
                    if col_idx in (1, 3, 5, 6, 7):
                        cell.alignment = Alignment(horizontal="center", vertical="center")

            # Atur lebar kolom
            col_widths = {
                "A": 8,
                "B": 25,
                "C": 10,
                "D": 40,
                "E": 18,
                "F": 18,
                "G": 15,
            }
            for col_letter, width in col_widths.items():
                ws.column_dimensions[col_letter].width = width

            # Simpan file
            wb.save(file_path)

            # Tampilkan dialog berhasil (Yes/No)
            self._show_export_success(file_path)

        except Exception as e:
            print(f"[ERROR] Gagal export Excel: {e}")
            messagebox.showerror("Error", f"Terjadi kesalahan: {e}")

    def export_pdf(self):
        """Export laporan ke file PDF menggunakan ReportLab.

        Membuat file Laporan_Perpustakaan_<tanggal-jam>.pdf berisi
        judul, tabel transaksi peminjaman, dan tanggal cetak. Menyimpan
        ke folder exports/pdf/ lalu menampilkan dialog berhasil.
        """
        try:
            # Pastikan data laporan sudah dimuat sesuai filter
            self.load_data()

            # Ambil data dari Treeview (mengikuti filter yang aktif)
            items = self.tree.get_children()
            if not items:
                messagebox.showwarning(
                    "Peringatan",
                    "Tidak ada data transaksi untuk diekspor.",
                )
                return

            # Folder tujuan dan nama file unik (tanggal & jam)
            folder_path = self._get_export_folder("pdf")
            timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
            file_name = f"Laporan_Perpustakaan_{timestamp}.pdf"
            file_path = os.path.join(folder_path, file_name)

            # Buat dokumen PDF ukuran A4 landscape agar tabel cukup lebar
            doc = SimpleDocTemplate(
                file_path,
                pagesize=landscape(A4),
                rightMargin=15 * mm,
                leftMargin=15 * mm,
                topMargin=15 * mm,
                bottomMargin=15 * mm,
            )

            elements = []

            # --- Styles ---
            styles = getSampleStyleSheet()

            # Gaya judul
            title_style = ParagraphStyle(
                "JudulStyle",
                parent=styles["Title"],
                fontName="Helvetica-Bold",
                fontSize=18,
                leading=22,
                alignment=1,  # center
                textColor=colors.HexColor("#2c3e50"),
                spaceAfter=10,
            )

            # Gaya tanggal cetak
            date_style = ParagraphStyle(
                "TanggalCetakStyle",
                parent=styles["Normal"],
                fontName="Helvetica",
                fontSize=10,
                alignment=1,
                textColor=colors.HexColor("#7f8c8d"),
                spaceAfter=15,
            )

            # --- Judul ---
            elements.append(Paragraph("Laporan Peminjaman Perpustakaan", title_style))

            # --- Tanggal cetak ---
            now_str = datetime.now().strftime("%d-%m-%Y %H:%M:%S")
            elements.append(Paragraph(f"Tanggal Cetak : {now_str}", date_style))

            elements.append(Spacer(1, 5 * mm))

            # --- Data tabel ---
            header = [
                "ID",
                "Nama Siswa",
                "Kelas",
                "Buku",
                "Tanggal Pinjam",
                "Tanggal Kembali",
                "Status",
            ]

            data = [header]
            for item in items:
                values = self.tree.item(item, "values")
                row = [str(v) if v is not None else "" for v in values]
                data.append(row)

            # Buat tabel dengan lebar kolom
            col_widths = [
                20 * mm,
                55 * mm,
                20 * mm,
                70 * mm,
                45 * mm,
                45 * mm,
                35 * mm,
            ]
            table = Table(data, colWidths=col_widths, repeatRows=1)

            # Style tabel
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2c3e50")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, 0), 10),
                        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                        ("FONTSIZE", (0, 1), (-1, -1), 9),
                        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                        ("ALIGN", (0, 1), (0, -1), "CENTER"),
                        ("ALIGN", (2, 1), (2, -1), "CENTER"),
                        ("ALIGN", (4, 1), (-1, -1), "CENTER"),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#bdc3c7")),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#ecf0f1")]),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )

            elements.append(table)

            # --- Build PDF ---
            doc.build(elements)

            # Tampilkan dialog berhasil (Yes/No)
            self._show_export_success(file_path)

        except Exception as e:
            print(f"[ERROR] Gagal export PDF: {e}")
            messagebox.showerror("Error", f"Terjadi kesalahan: {e}")

    def refresh_laporan(self):
        """Muat ulang data laporan."""
        try:
            self.load_data()
            messagebox.showinfo("Informasi", "Laporan berhasil di-refresh.")
        except Exception as e:
            print(f"[ERROR] Gagal refresh laporan: {e}")
            messagebox.showerror("Error", f"Terjadi kesalahan: {e}")

    def on_button_click(self, action):
        """Handler tombol."""
        if action == "tampilkan":
            self.tampilkan_laporan()
        elif action == "export_excel":
            self.export_excel()
        elif action == "export_pdf":
            self.export_pdf()
        elif action == "refresh":
            self.refresh_laporan()
        else:
            messagebox.showinfo(
                "Informasi",
                f"Fungsi '{action.capitalize()}' belum tersedia.",
            )
